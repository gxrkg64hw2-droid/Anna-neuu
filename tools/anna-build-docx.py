#!/usr/bin/env python
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/home/user/Anna-neuu/thesen/anna-postmigrantisch-theater/arbeit.md"
OUT = sys.argv[1]

doc = Document()

# --- Seitenränder ---
sec = doc.sections[0]
sec.left_margin, sec.right_margin = Cm(2.5), Cm(3.5)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)

# --- Normal-Stil: Times New Roman 12, 1,5-zeilig, Blocksatz ---
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_after = Pt(0)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Auto-Silbentrennung aktivieren
settings = doc.settings.element
for tag in ("w:autoHyphenation",):
    el = OxmlElement(tag); el.set(qn("w:val"), "true"); settings.append(el)

def style_heading(level, size, before, after):
    st = doc.styles[f"Heading {level}"]
    st.font.name = "Times New Roman"; st.font.size = Pt(size)
    st.font.bold = True; st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    st.paragraph_format.keep_with_next = True
style_heading(1,14,12,6); style_heading(2,12,8,4)

def add_runs(p, text):
    # **fett** unterstützen
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part == "": continue
        r = p.add_run(part); r.bold = (i % 2 == 1)

def page_break():
    doc.add_page_break()

# ---------------- Titelblatt ----------------
def center(text, size=12, bold=False, before=0, after=0):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    return p

center("[Universität / Hochschule]", 12, before=0, after=0)
center("[Institut / Fachbereich]", 12, after=0)
center("Seminar: [Seminartitel]", 12, after=0)
center("Leitung: [Titel Vorname Nachname der Seminarleitung]", 12, after=0)
center("[Semester, z. B. Sommersemester 2026]", 12, after=24)
for _ in range(3): doc.add_paragraph()
center("Kunst und Gesellschaftskritik", 17, bold=True, after=0)
center("Theater und Satire als Räume postmigrantischer Sichtbarkeit und Zugehörigkeit",
       14, bold=True, after=6)
center("Seminararbeit", 12, after=0)
for _ in range(6): doc.add_paragraph()
center("vorgelegt von", 12, after=0)
center("[Vor- und Nachname]", 12, bold=True, after=0)
center("Studienkennzahl: [xxx]   ·   Matrikelnummer: [xxx]", 12, after=0)
center("[E-Mail-Adresse]", 12, after=24)
center("Abgabedatum: 09. August 2026", 12, after=0)
page_break()

# ---------------- Inhaltsverzeichnis (Word-Feld) ----------------
h = doc.add_paragraph(); h.style = doc.styles["Heading 1"]; h.add_run("Inhaltsverzeichnis")
p = doc.add_paragraph()
run = p.add_run()
fldChar = OxmlElement('w:fldChar'); fldChar.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
instr.text = 'TOC \\o "1-2" \\h \\z \\u'
fldSep = OxmlElement('w:fldChar'); fldSep.set(qn('w:fldCharType'), 'separate')
t = OxmlElement('w:t'); t.text = "Rechtsklick > Feld aktualisieren (bzw. in Word: Verweise > Inhaltsverzeichnis aktualisieren)."
fldEnd = OxmlElement('w:fldChar'); fldEnd.set(qn('w:fldCharType'), 'end')
for e in (fldChar, instr, fldSep, t, fldEnd): run._r.append(e)
page_break()

# ---------------- Fließtext aus arbeit.md ----------------
lines = open(SRC, encoding="utf-8").read().splitlines()
started = False
for ln in lines:
    s = ln.rstrip()
    if s.startswith("# "):  # H1-Titel des Markdowns überspringen
        continue
    if s.startswith(">") or s.strip().startswith("<!--"):
        continue
    if s.strip() == "---":
        continue
    if s.startswith("### "):
        p = doc.add_paragraph(); p.style = doc.styles["Heading 2"]; add_runs(p, s[4:].strip()); started = True; continue
    if s.startswith("## "):
        title = s[3:].strip()
        p = doc.add_paragraph(); p.style = doc.styles["Heading 1"]; add_runs(p, title); started = True; continue
    if not s.strip():
        continue
    if not started:
        continue
    p = doc.add_paragraph(); add_runs(p, s.strip())

# ---------------- Literaturverzeichnis ----------------
doc.add_page_break()
p = doc.add_paragraph(); p.style = doc.styles["Heading 1"]; p.add_run("Literaturverzeichnis")
refs = [
 "Balme, Christopher B. (2022): Theaterwissenschaft und Postkolonialismus. Ein fachgeschichtlicher und persönlicher Rückblick. In: Sharifi, Azadeh/Skwirblies, Lisa (Hg.): Theaterwissenschaft postkolonial/dekolonial. Eine kritische Bestandsaufnahme. Bielefeld: transcript, S. 225–246.",
 "Carvalho, Wagner/Larsson, Fabian (2022): Ballhaus Naunynstraße. Theater einer offenen Gesellschaft. In: Sharifi, Azadeh/Skwirblies, Lisa (Hg.): Theaterwissenschaft postkolonial/dekolonial. Eine kritische Bestandsaufnahme. Bielefeld: transcript, S. 247–256.",
 "Hill, Marc/Yıldız, Erol (Hg.) (2018): Postmigrantische Visionen. Erfahrungen – Ideen – Reflexionen. Bielefeld: transcript.",
 "Langhoff, Shermin (2011): Die Herkunft spielt keine Rolle – „Postmigrantisches“ Theater im Ballhaus Naunynstraße. Interview, bpb.de, 10.03.2011. Online: https://www.bpb.de [Abruf: 06.08.2026].",
 "Sharifi, Azadeh/Skwirblies, Lisa (Hg.) (2022): Theaterwissenschaft postkolonial/dekolonial. Eine kritische Bestandsaufnahme. Bielefeld: transcript.",
 "Yıldız, Erol (2018): Vom methodologischen Nationalismus zu postmigrantischen Visionen. In: Hill, Marc/Yıldız, Erol (Hg.): Postmigrantische Visionen. Erfahrungen – Ideen – Reflexionen. Bielefeld: transcript, S. 43–62.",
 "Yıldız, Erol (2025): Postmigrantisch denken. Heimisch in einer globalisierten Gesellschaft. Bielefeld: transcript.",
]
for r in refs:
    par = doc.add_paragraph(r)
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.left_indent = Cm(0.75)
    par.paragraph_format.first_line_indent = Cm(-0.75)  # hängender Einzug

doc.save(OUT)
print("DOCX gespeichert:", OUT)
